from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
import base64
root=Path(__file__).resolve().parents[1]
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
blank=p.slide_layouts[6]
NAVY=RGBColor(17,44,66);BLUE=RGBColor(43,91,129);TEAL=RGBColor(21,129,128);MUTED=RGBColor(93,111,124);PALE=RGBColor(244,248,250);RED=RGBColor(185,67,70);GOLD=RGBColor(197,125,43)
def rect(sl,x,y,w,h,fill,border=None,r=0):
    shape=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=fill;shape.line.fill.background() if border is None else None
    if border:shape.line.color.rgb=border
    return shape
def tx(sl,x,y,w,h,value,size=16,color=NAVY,bold=False,align=None):
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));t=sh.text_frame;t.clear();t.word_wrap=True
    t.margin_left=Pt(1);t.margin_top=Pt(1);t.margin_bottom=Pt(1);t.margin_right=Pt(1)
    paragraph=t.paragraphs[0];paragraph.text=value;paragraph.font.name='Aptos';paragraph.font.size=Pt(size);paragraph.font.bold=bold;paragraph.font.color.rgb=color
    if align:paragraph.alignment=align
    return sh
def base(title,number):
    s=p.slides.add_slide(blank);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(255,255,255)
    rect(s,0,0,13.333,.12,TEAL);tx(s,.65,.27,5.0,.31,'M I G H T O R A   /   R A A I D D',11,TEAL,True)
    tx(s,.65,.82,12, .7,title,29,NAVY,True);rect(s,.65,1.62,12,.025,RGBColor(220,229,236))
    tx(s,.65,7.15,10,.2,'Mightora RAAIDD • Client-side report • Generated from imported data',9,MUTED)
    tx(s,12.1,7.14,.5,.2,str(number),9,MUTED,align=PP_ALIGN.RIGHT)
    return s
# Cover
s=p.slides.add_slide(blank);s.background.fill.solid();s.background.fill.fore_color.rgb=NAVY
rect(s,.55,.6,.09,5.9,TEAL);tx(s,.96,.76,11,.35,'M I G H T O R A  /  G O V E R N A N C E',13,RGBColor(126,218,210),True)
tx(s,.96,1.76,11,1.45,'{{S1_TITLE}}',38,RGBColor(255,255,255),True)
tx(s,.98,3.34,10,.8,'{{S1_SUBTITLE}}',19,RGBColor(210,227,240))
tx(s,.98,5.68,8,.4,'{{S1_DATE}}',14,RGBColor(210,227,240))
tx(s,.98,6.73,8,.27,'Prepared locally. No server processing required.',10,RGBColor(210,227,240))
# Summary
s=base('Executive overview',2)
for i,(name,x,clr) in enumerate([('TOTAL',.65,BLUE),('OPEN',3.15,TEAL),('HIGH RISKS',5.65,RED),('OVERDUE',8.15,GOLD),('QUALITY',10.65,BLUE)]):
 rect(s,x,2.0,2.12,1.55,PALE);rect(s,x,2.0,2.12,.08,clr);tx(s,x+.14,2.22,1.7,.26,name,10,MUTED,True);tx(s,x+.14,2.58,1.8,.65,f'{{{{S2_KPI_{i}}}}}',31,clr,True)
tx(s,.7,4.16,11.6,.46,'{{S2_HEADLINE}}',21,NAVY,True)
tx(s,.7,4.91,11.8,1.0,'{{S2_NOTE}}',18,MUTED)
# Heatmap
s=base('Risk exposure matrix',3)
tx(s,.68,1.8,10,.34,'Active risks by likelihood × impact (1–5)',16,MUTED)
for j in range(5):tx(s,2.3+j*.88,2.28,.7,.3,str(j+1),12,MUTED,True,PP_ALIGN.CENTER)
for i in range(5):
 y=2.67+i*.69; likelihood=5-i;tx(s,1.72,y+.15,.4,.25,str(likelihood),12,MUTED,True,PP_ALIGN.CENTER)
 for j in range(5):
  impact=j+1;sc=likelihood*impact; col=RGBColor(220,241,230) if sc<=5 else RGBColor(249,232,189) if sc<=11 else RGBColor(248,204,158) if sc<=19 else RGBColor(244,169,160)
  x=2.3+j*.88;rect(s,x,y,.75,.57,col);tx(s,x+.07,y+.11,.6,.31,f'{{{{H_{likelihood}_{impact}}}}}',13,NAVY,True,PP_ALIGN.CENTER)
tx(s,7.65,2.65,4.5,.44,'{{S3_COUNT}}',23,TEAL,True)
tx(s,7.65,3.3,4.2,1.8,'Heatmap counts show active risks where both numeric likelihood and impact are recorded. Blank scores are not treated as zero.',15,MUTED)
# Attention
s=base('Items requiring attention',4)
for i in range(5):
 y=1.91+i*.92;rect(s,.67,y,12,.77,PALE);rect(s,.67,y,.08,.77,GOLD if i<3 else TEAL)
 tx(s,.93,y+.08,1.05,.25,f'{{{{S4_ID_{i}}}}}',11,TEAL,True);tx(s,2.12,y+.06,9.85,.32,f'{{{{S4_TITLE_{i}}}}}',14,NAVY,True);tx(s,2.12,y+.41,9.5,.2,f'{{{{S4_DETAIL_{i}}}}}',10,MUTED)
# Decisions/actions
s=base('Decisions and overdue actions',5)
tx(s,.72,1.86,6,.36,'DECISIONS REQUIRING ATTENTION',15,TEAL,True)
tx(s,7.0,1.86,6,.36,'OVERDUE ACTIONS',15,GOLD,True)
for i in range(4):
 y=2.4+i*.89;rect(s,.69,y,5.95,.71,PALE);rect(s,6.98,y,5.65,.71,PALE)
 tx(s,.84,y+.1,5.65,.56,f'{{{{S5_DEC_{i}}}}}',12,NAVY)
 tx(s,7.12,y+.1,5.3,.56,f'{{{{S5_ACT_{i}}}}}',12,NAVY)
# Assurance
s=base('Assurance findings and next steps',6)
tx(s,.71,1.9,11,.4,'{{S6_HEADLINE}}',19,NAVY,True)
for i in range(4):
 y=2.65+i*.91;rect(s,.69,y,11.99,.72,PALE);rect(s,.69,y,.06,.72,TEAL)
 tx(s,.89,y+.11,11.2,.48,f'{{{{S6_FIND_{i}}}}}',13,NAVY)
out=root/'src/assets/report-template.pptx';p.save(out)
b64=base64.b64encode(out.read_bytes()).decode('ascii')
(root/'src/pptx-template.js').write_text('const PPTX_TEMPLATE_BASE64 = '+repr(b64)+';\n')
print(out.stat().st_size,'bytes',len(b64),'base64')