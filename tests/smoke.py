"""Browser end-to-end tests. Runs in Chromium without network or a backend."""
from pathlib import Path
from playwright.sync_api import sync_playwright
import tempfile, zipfile, re, json, os, subprocess, time, shutil
ROOT=Path(__file__).resolve().parents[1]
DIST=ROOT/'dist'
assert (DIST/'raaidd-offline.html').exists(),'Run python scripts/build.py first'

def check_office(file, expected):
    with zipfile.ZipFile(file) as z:
        names=z.namelist()
        assert '[Content_Types].xml' in names, (file,names)
        assert expected in names, (file,expected)
        for name in names:
            if name.endswith('.xml'):
                from xml.etree import ElementTree as ET
                try:ET.fromstring(z.read(name))
                except Exception as e:raise AssertionError(f'{file.name}: malformed {name}: {e}')

with tempfile.TemporaryDirectory() as td, sync_playwright() as pw:
    tmp=Path(td)
    browser=pw.chromium.launch(headless=True, executable_path='/usr/bin/chromium' if Path('/usr/bin/chromium').exists() else None,args=['--no-sandbox'])
    page=browser.new_page(accept_downloads=True, viewport={'width':1350,'height':850});page.set_default_timeout(7000)
    errors=[];requests=[]
    page.on('pageerror',lambda error:errors.append(str(error)))
    page.on('request',lambda req:requests.append(req.url) if req.url.startswith('http') else None)
    page.set_content((DIST/'raaidd-offline.html').read_text(),wait_until='load')
    assert page.locator('#offline-banner').is_visible(),'offline banner missing'
    assert page.locator('h1').inner_text().startswith('Your RAID')
    page.locator('#demo').click()
    assert len(page.locator('.record-list-row').all())==10
    assert 'Loaded fictional sample data' in page.locator('#status-line').inner_text()
    # Relationship graph must render existing links and allow navigation to linked records.
    page.locator('[data-tab="relationships"]').click()
    assert page.locator('#tab-relationships').is_visible()
    assert page.locator('.graph-node').count() >= 5
    assert page.locator('.linked-row').count() >= 5
    assert 'unresolved reference' in page.locator('#graph-summary').inner_text()
    page.locator('#graph-type').select_option('Risk')
    assert page.locator('.graph-node').count() >= 2, 'type filtering hid connected neighbours'
    page.locator('#graph-type').select_option('')
    page.locator('#graph-focus').select_option(label='R-001 · Integration supplier may miss committed delivery')
    assert page.locator('.graph-node').count() >= 2
    page.locator('#graph-focus').select_option('')
    page.locator('.graph-node').first.click()
    assert page.locator('#tab-register').is_visible(), 'graph node should open record editor'
    # New relationships are added and removed in the register without losing other data.
    before=page.locator('.detail-link-row').count()
    page.locator('#relation-target').select_option('DEC-001')
    page.locator('#add-relation').click()
    assert page.locator('.detail-link-row').count()==before+1
    page.locator('[data-tab="relationships"]').click()
    assert page.locator('.linked-row').count()>=6
    page.locator('[data-tab="register"]').click()
    page.locator('.remove-link').last.click()
    assert page.locator('.detail-link-row').count()==before
    # Evidence section is positioned right before the footer and uses real workflow URLs.
    assert page.locator('mightora-footer').count()==1
    assert page.locator('mightora-footer').evaluate('(el) => el.previousElementSibling.classList.contains("trust")')
    assert page.locator('.badges a').count()==5
    assert page.locator('.badges img').count()==0, 'offline edition should not fetch workflow badges'
    source_html=(ROOT/'index.html').read_text()
    assert source_html.count('class="proof-img"')==3, 'hosted edition should use live workflow images'
    assert 'actions/workflows/codeql.yml/badge.svg' in source_html
    assert 'actions/workflows/ci.yml/badge.svg' in source_html
    assert 'actions/workflows/pages.yml/badge.svg' in source_html
    assert not page.locator('html').evaluate('(el) => el.outerHTML.includes("github.com/mightora/raid.mightora.io/actions/workflows/ci.yml/badge.svg")')
    assert page.locator('mightora-header').evaluate('(el) => getComputedStyle(el).position')=='sticky'
    # edit and review
    page.locator('[data-edit="title"]').fill('Supplier API delay — testing edition')
    page.locator('#save-record').click()
    assert 'Supplier API delay' in page.locator('.record-panel h4').first.inner_text()
    page.locator('#comment-text').fill('Testing local reviewer comments in the smoke test.')
    page.locator('#add-comment').click()
    assert page.locator('.comment').count()>=2
    page.locator('[data-tab="insights"]').click()
    assert page.locator('.heatcell').count()==25
    page.locator('#take-snapshot').click()
    assert page.locator('.snaprow').count()==1
    page.locator('[data-tab="reports"]').click()
    outputs={}
    for id,ext in [('export-enhanced','xlsx'),('export-word','docx'),('export-slides','pptx'),('export-project','json')]:
        with page.expect_download(timeout=25000) as info:page.locator('#'+id).click()
        item=info.value
        save=tmp/f'generated.{ext}';item.save_as(save);outputs[ext]=save
        assert save.stat().st_size>800,(ext,save.stat().st_size)
    check_office(outputs['xlsx'],'xl/workbook.xml')
    check_office(outputs['docx'],'word/document.xml')
    check_office(outputs['pptx'],'ppt/presentation.xml')
    project=json.loads(outputs['json'].read_text())
    assert len(project['records'])==10
    assert project['snapshots'] and project['records'][0]['_comments']
    # Reimport enhanced output; it has each category as a separate worksheet
    page.locator('[data-tab="import"]').click()
    print('reimport xlsx',flush=True);page.locator('#file-input').set_input_files(outputs['xlsx']);page.wait_for_timeout(750);print('status',page.locator('#status-line').inner_text(),flush=True)
    assert page.locator('#import-editor').is_visible(), page.locator('#status-line').inner_text()
    print('sheets',page.locator('#sheet-select option').count(),flush=True);assert page.locator('#sheet-select option').count()>=7
    page.locator('#load-register').click()
    assert len(page.locator('.record-list-row').all())>=5, page.locator('#status-line').inner_text()
    # round-trip original-layout export and parse workbook contents
    page.locator('[data-tab="reports"]').click()
    with page.expect_download() as info:page.locator('#export-original').click()
    original=tmp/'original.xlsx';info.value.save_as(original)
    check_office(original,'xl/workbook.xml')
    # load project file and preserve reviewer comments
    page.locator('#file-input').set_input_files(outputs['json']);page.wait_for_timeout(600);print('project status',page.locator('#status-line').inner_text(), 'comments', page.locator('.comment').count(),flush=True)
    assert page.locator('.comment').count()>=2
    bad=json.loads(outputs['json'].read_text());bad['records'][0]['_relationships'].append({'type':'Blocks','target':'MISSING-EXAMPLE'})
    invalid=tmp/'bad-relations.raaidd.json';invalid.write_text(json.dumps(bad))
    page.locator('#file-input').set_input_files(invalid)
    page.wait_for_timeout(300)
    page.locator('[data-tab="relationships"]').click()
    assert page.locator('#graph-summary span').nth(2).locator('b').inner_text()=='1'
    assert page.locator('.linked-row.has-warning').count()==1
    page.locator('#file-input').set_input_files(outputs['json'])
    # Export Word & PowerPoint parsable by office libs
    from docx import Document
    document=Document(outputs['docx'])
    assert 'RAAIDD Governance' in '\n'.join(p.text for p in document.paragraphs)
    from pptx import Presentation
    prs=Presentation(outputs['pptx'])
    assert len(prs.slides)==6
    texts=[shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame]
    assert any('Supplier API delay' in t for t in texts)
    assert sum(len(slide.shapes) for slide in prs.slides)>70
    assert not requests, 'Offline app made an HTTP request: '+repr(requests)
    assert not errors,'Browser errors: '+repr(errors)
    print('OFFLINE PASS: demo, review, comments, risk matrix, snapshots, XLSX / DOCX / PPTX / project export, XLSX re-import, privacy')
    print('Export sizes:',{k:v.stat().st_size for k,v in outputs.items()})

    # Excel header-offset/mapping and source-layout preservation test.
    def test_workbook(modified=False):
        return page.evaluate('''async (modified) => { const rows=[['Project RAID - cover'],['Guidance: header begins at row four'],[],['Ref','Description','Assigned To','RAG','Due','Notes'],['R-900','Risk of losing source field', 'Example owner', modified?'Closed':'Open','2026-12-12','Keep original'],...(modified?[['R-901','New supplier problem','Example owner','Open','2026-12-20','Other text']]:[])];const blob=Spreadsheet.newWorkbook({'Project RAID':rows});return [...new Uint8Array(await blob.arrayBuffer())]; }''',modified)
    fixture=tmp/'offset-header.xlsx';fixture.write_bytes(bytes(test_workbook()))
    updated=tmp/'updated-source.xlsx';updated.write_bytes(bytes(test_workbook(True)))
    page.locator('#file-input').set_input_files(fixture)
    page.wait_for_timeout(450)
    assert page.locator('#header-row').input_value()=='4', 'header-row autodetection failed'
    page.locator('[data-export-field="mitigation"]').check()
    page.locator('#load-register').click()
    assert page.locator('.record-list-row').count()==1
    page.locator('[data-edit="status"]').fill('In progress')
    page.locator('[data-edit="mitigation"]').fill('Agree supplier test evidence')
    page.locator('#save-record').click()
    page.locator('[data-tab="reports"]').click()
    with page.expect_download() as info:page.locator('#export-original').click()
    preserved=tmp/'source-preserved.xlsx';info.value.save_as(preserved)
    with zipfile.ZipFile(preserved) as z:
        content=z.read('xl/worksheets/sheet1.xml').decode()
        assert 'Keep original' in content and 'Agree supplier test evidence' in content
        assert 'In progress' in content and 'Mitigation / response' in content
        assert 'Project RAID - cover' in content
    page.locator('[data-tab="register"]').click()
    page.locator('#compare-input').set_input_files(updated)
    page.wait_for_timeout(450)
    assert page.locator('#compare-results').is_visible(),page.locator('#status-line').inner_text()
    assert '1 new' in page.locator('.comparison-summary').inner_text()
    assert '1 changed' in page.locator('.comparison-summary').inner_text()
    print('OFFSET-HEADER PASS: custom header row, unmapped source field, optional column, safe round-trip, comparison')
    # CSV additional-column test. Spreadsheet formulas in source remain literal strings.
    csv=tmp/'csv-test.csv';csv.write_text('Ref,Description,Assigned To,RAG,Notes\nR-2,CSV example,Owner A,Open,Original note\n')
    page.locator('#file-input').set_input_files(csv);page.wait_for_timeout(350)
    page.locator('[data-export-field="mitigation"]').check()
    page.locator('#load-register').click()
    page.locator('[data-edit="mitigation"]').fill('=HYPERLINK("https://example.invalid")')
    page.locator('#save-record').click()
    page.locator('[data-tab="reports"]').click()
    with page.expect_download() as info:page.locator('#export-original').click()
    csv_out=tmp/'csv-updated.csv';info.value.save_as(csv_out)
    out=csv_out.read_text(encoding='utf-8-sig')
    assert 'Original note' in out and 'Mitigation / response' in out and '=HYPERLINK' in out
    print('CSV PASS: source fields preserved and optional fields appended')
    # Report sections control native PowerPoint slide inclusion.
    page.locator('.section-picks input[value="risks"]').uncheck()
    page.locator('.section-picks input[value="evidence"]').uncheck()
    with page.expect_download() as info:page.locator('#export-slides').click()
    short_pptx=tmp/'short.pptx';info.value.save_as(short_pptx)
    assert len(Presentation(short_pptx).slides)==3
    print('REPORT SECTIONS PASS: reduced PowerPoint contains 3 editable slides')

    
    # Keep example files so end users can inspect valid editable Office reports.
    for ext,path in outputs.items():
        shutil.copy2(path,ROOT/'docs'/('example-governance.'+('raaidd.json' if ext=='json' else ext)))
    shutil.copy2(fixture,ROOT/'docs'/'example-offset-header.xlsx')
    # save screenshots of tested UI for visual verification (not distributed)
    page.locator('[data-tab="insights"]').click()
    page.screenshot(path=str(ROOT/'docs/preview-insights.png'),full_page=True)
    page.locator('[data-tab="register"]').click()
    page.screenshot(path=str(ROOT/'docs/preview-register.png'),full_page=True)
    browser.close()